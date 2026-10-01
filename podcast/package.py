"""Stages: transcript cleaning, episode copy, and the episode package
(``.docx`` in the owner's episode-document structure + website HTML).

The website HTML mirrors the markup of the live "Inspiring conversations"
episode pages (two Squarespace text blocks around a video block): H3 title +
intro, then H4 key insights (bold group line + "Lead: text" bullets) and H4
full transcript, a bold name paragraph before each speaker turn.
"""

from __future__ import annotations

import hashlib
import html
import json
import logging
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from podcast.clips import clip_length, load_clips
from podcast.episode import Episode, work_dir
from podcast.hub import ask, ask_json
from podcast.metrics import StageRecord
from podcast.review_state import kept_clips, load_review
from podcast.transcribe import fmt_ts, load_turns

logger = logging.getLogger("podcast.package")

CLEAN_FILE = "transcript/clean.json"
EPISODE_COPY_FILE = "episode_copy.json"
CHUNK_CHARS = 7000
MIN_KEPT_RATIO = 0.55


# ── clean ───────────────────────────────────────────────────────────────

CLEAN_SYSTEM = "You are a careful transcript editor. You never summarise or add content."

CLEAN_PROMPT = """Edit this part of an interview transcript for reading.

- Keep every paragraph, in order, each starting with its speaker label exactly as given ("Name: ").
- Fix words that were clearly misheard, using the context.
- Remove filler (um, uh, "you know" and "like" used as filler), false starts, stutters and accidental repetitions.
- Fix punctuation and capitalisation. You may split a very long paragraph; continuation paragraphs of the same speaker start without a label.
- Do not rephrase, shorten or summarise beyond that. Keep the speaker's own words and tone.

Reply with the edited paragraphs only, separated by blank lines.

{text}
"""


def chunk_paragraphs(paragraphs: list[str], limit: int = CHUNK_CHARS) -> list[list[str]]:
    chunks: list[list[str]] = [[]]
    size = 0
    for p in paragraphs:
        if chunks[-1] and size + len(p) > limit:
            chunks.append([])
            size = 0
        chunks[-1].append(p)
        size += len(p)
    return [c for c in chunks if c]


def parse_clean(text: str, labels: tuple[str, ...]) -> list[dict]:
    """``Label: text`` paragraphs → ``{speaker, text}``; unlabeled ones continue the speaker."""
    out: list[dict] = []
    label_re = re.compile(rf"^\**({'|'.join(map(re.escape, labels))})\**:\s*\**\s*(.*)$")
    for block in re.split(r"\n\s*\n", text.strip()):
        block = " ".join(block.split())
        if not block:
            continue
        m = label_re.match(block)
        if m:
            out.append({"speaker": m.group(1), "text": m.group(2).strip()})
        elif out:
            out.append({"speaker": None, "text": block})
    return out


def _accept(raw: list[dict], cleaned: list[dict]) -> tuple[bool, str]:
    """Reject an edit that lost text or speaker labels (a truncated or off-task reply)."""
    kept = sum(len(p["text"]) for p in cleaned) / max(1, sum(len(p["text"]) for p in raw))
    labelled = sum(1 for p in cleaned if p["speaker"])
    ok = kept >= MIN_KEPT_RATIO and labelled >= 0.8 * len(raw)
    return ok, f"kept {100 * kept:.0f}%, {labelled}/{len(raw)} labels"


def _clean_chunk(cfg: dict, rec: StageRecord, chunk: list[str], labels: tuple[str, ...],
                 cache_dir: Path, attempts: int = 2) -> list[dict]:
    """Edit one chunk; an accepted edit is cached by source hash so a re-run only redoes the rest."""
    source = "\n\n".join(chunk)
    raw = parse_clean(source, labels)
    cached = cache_dir / f"{hashlib.sha1(source.encode('utf-8')).hexdigest()[:16]}.json"
    if cached.exists():
        return json.loads(cached.read_text(encoding="utf-8"))
    for attempt in range(1, attempts + 1):
        try:
            reply = ask(cfg, rec, "clean", CLEAN_PROMPT.format(text=source), system=CLEAN_SYSTEM,
                        max_tokens=8000)
        except Exception as exc:  # noqa: BLE001 — keep the raw chunk rather than lose the episode
            logger.error("❌ clean chunk failed, keeping raw text: %s", exc)
            return raw
        cleaned = parse_clean(reply, labels)
        ok, why = _accept(raw, cleaned)
        if ok:
            cached.write_text(json.dumps(cleaned, ensure_ascii=False), encoding="utf-8")
            return cleaned
        logger.warning("⚠️ clean chunk rejected on attempt %d (%s); reply starts: %r",
                       attempt, why, reply[:300])
    logger.warning("⚠️ keeping the raw text for a chunk starting %r", source[:80])
    return raw


def run_clean(ep: Episode, cfg: dict, rec: StageRecord) -> Path:
    labels = (ep.label("guest"), ep.label("host"))
    paragraphs = [f"{ep.label(t['speaker'])}: {t['text']}" for t in load_turns(ep)]
    chunks = chunk_paragraphs(paragraphs)
    logger.info("ℹ️ cleaning %d paragraphs in %d chunks", len(paragraphs), len(chunks))
    cache_dir = work_dir(cfg, ep) / "clean_cache"
    cache_dir.mkdir(exist_ok=True)
    with ThreadPoolExecutor(max_workers=3) as pool:
        results = list(pool.map(lambda c: _clean_chunk(cfg, rec, c, labels, cache_dir), chunks))
    cleaned = [p for part in results for p in part]
    (ep.package / CLEAN_FILE).write_text(json.dumps(cleaned, indent=1, ensure_ascii=False), encoding="utf-8")
    md = [f"# {ep.base_name} - transcript\n",
          "_Generated locally from the recording and lightly edited for readability._\n"]
    txt = []
    for p in cleaned:
        md.append(f"**{p['speaker']}:** {p['text']}\n" if p["speaker"] else f"{p['text']}\n")
        txt.append(f"{p['speaker']}: {p['text']}" if p["speaker"] else p["text"])
    (ep.package / f"{ep.base_name} - transcript.txt").write_text("\n\n".join(txt) + "\n", encoding="utf-8")
    path = ep.package / f"{ep.base_name} - transcript.md"
    path.write_text("\n".join(md), encoding="utf-8")
    logger.info("✅ cleaned transcript: %d paragraphs", len(cleaned))
    return path


def load_clean(ep: Episode) -> list[dict]:
    return json.loads((ep.package / CLEAN_FILE).read_text(encoding="utf-8"))


# ── episode copy ────────────────────────────────────────────────────────

EPISODE_PROMPT = """Below is the cleaned transcript of a conversation between {guest} and {host}. {roles}

Write the episode copy, in plain warm English, no hashtags, no em dashes. Reply with JSON only:
{{
 "youtube_title": "<max 70 characters, names allowed>",
 "youtube_hook": "<one question that opens the description>",
 "youtube_intro": "<one paragraph: who is talking and what the conversation covers>",
 "listen_bullets": ["<5 short bullets: what you get by listening>"],
 "website_title": "<the page title, like 'Transform your organization's execution and results, with {guest_first}'>",
 "website_intro": "<two short paragraphs introducing the conversation>",
 "insights": [{{"group": "<topic>", "bullets": [{{"lead": "<2-4 words>", "text": "<one sentence>"}}]}}],
 "thumbnail_title": "<the episode promise in at most 6 words>",
 "card_blurb": "<at most 45 words: who {guest_first} is and what we talk about>"
}}
Use 3 to 5 insight groups with 2 to 4 bullets each.

Transcript:
{transcript}
"""


def run_episode_copy(ep: Episode, cfg: dict, rec: StageRecord) -> dict:
    roles = (f"{ep.guest_first} interviews {ep.host_first}; {ep.host_first} is the guest of this conversation."
             if ep.host_is_interviewee else f"{ep.host_first} interviews {ep.guest_first}.")
    transcript = "\n\n".join(f"{p['speaker']}: {p['text']}" if p["speaker"] else p["text"]
                             for p in load_clean(ep))
    data = ask_json(cfg, rec, "episode_copy", EPISODE_PROMPT.format(
        guest=ep.guest_display, host=ep.host_name, guest_first=ep.guest_first,
        roles=roles, transcript=transcript), max_tokens=6000)
    (ep.package / EPISODE_COPY_FILE).write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
    logger.info("✅ episode copy: %s", data.get("youtube_title"))
    return data


def load_episode_copy(ep: Episode) -> dict:
    return json.loads((ep.package / EPISODE_COPY_FILE).read_text(encoding="utf-8"))


# ── website HTML ────────────────────────────────────────────────────────

_PRE = ' style="white-space:pre-wrap;"'


def _p(inner: str) -> str:
    return f'<p class=""{_PRE}>{inner}</p>'


def _link_first(text: str, name: str, href: str) -> str:
    """Escape ``text`` and link the first mention of ``name`` (the site links the guest's name)."""
    i = text.find(name)
    if i < 0:
        return html.escape(text)
    return (html.escape(text[:i]) + f'<a href="{html.escape(href)}" target="_blank">{html.escape(name)}</a>'
            + html.escape(text[i + len(name):]))


def website_html(ep: Episode, copy: dict, transcript: list[dict]) -> str:
    """The episode page's two text blocks, as on the live site, wrapped in a preview page.

    Block 1 is the H3 title and the intro (first mention of the guest linked
    to their LinkedIn); the YouTube video block sits between the two; block 2
    is the key insights and the full transcript, where each speaker turn is a
    bold name paragraph followed by its text paragraphs.
    """
    e = html.escape
    intro = [p for p in re.split(r"\n\s*\n", copy["website_intro"].strip()) if p]
    name = next((n for n in (ep.guest_display, ep.guest) if n in copy["website_intro"]), ep.guest)
    block1 = [f"<h3{_PRE}>{e(copy['website_title'])}</h3>"]
    block1 += [_p(_link_first(p, name, "#guest-linkedin") if i == 0 else e(p)) for i, p in enumerate(intro)]

    block2 = [f"<h4{_PRE}><strong>Key insights from the conversation</strong></h4>"]
    for group in copy.get("insights", []):
        block2.append(_p(f"<strong>{e(group['group'])}</strong>"))
        items = "".join(f"<li>{_p(e(b['lead'] + ': ' + b['text']))}</li>" for b in group.get("bullets", []))
        block2.append(f'<ul data-rte-list="default">{items}</ul>')
    block2.append(f"<h4{_PRE}><strong>Full transcript</strong></h4>")
    note = "Here is the entire conversation transcript, edited for clarity and conciseness."
    if ep.youtube_url:
        note += f' Here\'s also the link to the conversation on <a href="{e(ep.youtube_url)}" target="_blank">YouTube</a>.'
    block2.append(_p(note))
    for p in transcript:
        if p["speaker"]:
            block2.append(_p(f"<strong>{e(p['speaker'])}</strong>"))
        block2.append(_p(e(p["text"])))
    # episode.json "links": [{"label", "url"}]; until filled, a LinkedIn placeholder
    links = ep.extra.get("links") or [{"label": "LinkedIn", "url": "#guest-linkedin"}]
    block2.append(f"<h4{_PRE}><strong>Where to find {e(ep.guest_first)} and "
                  f"{e(ep.guest_pronoun_possessive)} work</strong></h4>")
    anchors = (f'<a href="{e(link["url"])}" target="_blank">{e(link["label"])}</a>' for link in links)
    block2.append(f'<ul data-rte-list="default">{"".join(f"<li>{_p(a)}</li>" for a in anchors)}</ul>')

    one, two = "\n".join(block1), "\n".join(block2)
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(copy['website_title'])}</title>
<style>
  body {{ background: #fff; color: #000; font-family: Poppins, Arial, sans-serif; max-width: 760px;
         margin: 0 auto; padding: 24px 16px; line-height: 1.6; }}
  h3 {{ font-size: 1.6rem; }} h4 {{ font-size: 1.2rem; margin-top: 2rem; }}
  a {{ color: #000; text-decoration-color: #FFCC00; }}
  .video {{ border: 2px dashed #999; padding: 24px; text-align: center; color: #666; margin: 24px 0; }}
</style>
</head>
<body>
<!-- Squarespace, new page under /inspiring-conversations/{e(ep.website_slug or "<slug>")}:
     paste BLOCK 1 into a text block, add a Video block with the YouTube link,
     then paste BLOCK 2 into a second text block. Replace #guest-linkedin first. -->
<!-- BLOCK 1 START -->
<div class="sqs-html-content">
{one}
</div>
<!-- BLOCK 1 END -->
<div class="video">YouTube video block{": " + e(ep.youtube_url) if ep.youtube_url else ""}</div>
<!-- BLOCK 2 START -->
<div class="sqs-html-content">
{two}
</div>
<!-- BLOCK 2 END -->
</body>
</html>
"""


# ── episode document ────────────────────────────────────────────────────

def _thank_you(ep: Episode, n_clips: int) -> list[str]:
    """Paragraphs of the owner's thank-you note; links stay placeholders until published."""
    slug = ep.website_slug or "(slug)"
    return [
        f"Hi {ep.guest_first}!",
        "I hope you are having a great weekend.",
        "I am happy to share with you that our conversation is live on my website:",
        f"https://www.robertoferraro.net/inspiring-conversations/{slug}",
        "YouTube link",
        ep.youtube_url or "(to add)",
        "Plus on Spotify and on all the main podcast platforms.",
        "(to add)",
        f"I also edited {n_clips} clips, each in square (1:1) and vertical (9:16) format, with subtitles. "
        "It's all here in OneDrive with the rest:",
        "(folder link)",
        "Of course, again feel free to use and share whatever you like. It's yours!",
        "Thanks again for the great time.",
        "Kind regards and let's keep in touch,",
        ep.host_first,
    ]


def build_docx(ep: Episode, copy: dict, clips: list[dict], path: Path) -> Path:
    from docx import Document  # noqa: PLC0415 — only this stage needs python-docx

    doc = Document()
    doc.add_heading(ep.base_name, 0)

    doc.add_heading("Youtube", 1)
    doc.add_heading("Title", 2)
    doc.add_paragraph(copy["youtube_title"])
    doc.add_heading("Text", 2)
    doc.add_paragraph(copy["youtube_hook"])
    doc.add_paragraph(copy["youtube_intro"])
    doc.add_paragraph("📌 What will you get by listening?")
    for bullet in copy.get("listen_bullets", []):
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph(f"📌 Where to find {ep.guest_first}: (link)")

    doc.add_heading("Clips", 1)
    for c in clips:
        doc.add_heading(f"{c['number']}. {c['title']}", 2)
        files = c.get("file", "")
        doc.add_paragraph(f"{fmt_ts(c['start'])}-{fmt_ts(c['end'])} ({clip_length(c):.0f} s) · "
                          f"clips/1x1/{files}.mp4 · clips/9x16/{files}.mp4 · clips/covers/{files}.png")
        doc.add_paragraph().add_run("Instagram").bold = True
        doc.add_paragraph(c.get("instagram", ""))
        doc.add_paragraph().add_run("LinkedIn").bold = True
        for para in c.get("linkedin", "").split("\n\n"):
            doc.add_paragraph(para)

    doc.add_heading("Thumbnails", 1)
    doc.add_paragraph(f"Episode: {copy.get('thumbnail_title', '').upper()}")
    for c in clips:
        doc.add_paragraph(c["title"].upper(), style="List Bullet")

    doc.add_heading("Website text", 1)
    doc.add_heading("Title", 2)
    doc.add_paragraph(copy["website_title"])
    doc.add_heading("Intro", 2)
    for para in re.split(r"\n\s*\n", copy["website_intro"].strip()):
        doc.add_paragraph(para)
    doc.add_heading("Key insights", 2)
    for group in copy.get("insights", []):
        doc.add_paragraph().add_run(group["group"]).bold = True
        for b in group.get("bullets", []):
            item = doc.add_paragraph(style="List Bullet")
            item.add_run(f"{b['lead']}: ").bold = True
            item.add_run(b["text"])

    doc.add_heading("Links", 1)
    slug = ep.website_slug or "(slug)"
    for label, value in (("Own page", f"https://www.robertoferraro.net/inspiring-conversations/{slug}"),
                         ("Youtube", ep.youtube_url or "(to add)"),
                         (f"{ep.guest_first} LinkedIn", "(to add)"),
                         ("Spotify", "(to add)"), ("Substack", "(to add)"),
                         ("Episode package", str(ep.package))):
        doc.add_paragraph(f"{label}: {value}")

    doc.add_heading("thank you (en)", 1)
    for para in _thank_you(ep, len(clips)):
        doc.add_paragraph(para)
    doc.save(str(path))
    return path


def run_package(ep: Episode, cfg: dict, rec: StageRecord) -> Path:
    copy = load_episode_copy(ep)
    clips = kept_clips(load_clips(ep), load_review(ep))
    html_path = ep.package / f"{ep.base_name} - website.html"
    html_path.write_text(website_html(ep, copy, load_clean(ep)), encoding="utf-8")
    path = build_docx(ep, copy, clips, ep.package / f"{ep.base_name}.docx")
    logger.info("✅ package: %s, %s", path.name, html_path.name)
    return path
