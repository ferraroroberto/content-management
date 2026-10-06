"""Deterministic checks that stop a demo before it renders (issue #360).

Each check returns ``Result``s whose state is ``pass``, ``fail`` or
``unknown``; anything but ``pass`` stops the run. They exist because each one
caught a real defect in the hand-made reference video only by eye:

- ``overrun``: a sped-up clip ran past its beat into the next app screen;
- ``state_timeline``: a timer legend was synced to assumed, not measured, times;
- ``privacy``: made-up names shared a surname with real participants;
- ``licence``: copyrighted music must never reach a public cut.
"""

from __future__ import annotations

import csv
import json
import logging
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterable, Optional

import numpy as np

from demo_video.media import Result
from demo_video.storyboard import Cut, Demo, WindowStates, clip_files, load_marks, media_root
from podcast.media import ffmpeg_pipe, probe

logger = logging.getLogger("demo_video.checks")

# ----------------------------------------------------------------------------- overrun


def _played(scene: Any) -> list[tuple[str, dict, float]]:
    """``(where, clip ref, seconds of scene time it plays)`` for every played clip in a scene."""
    raw = scene.model_dump(mode="json")
    out: list[tuple[str, dict, float]] = []
    if isinstance(scene, WindowStates):
        used = 0.0
        for i, seg in enumerate(raw["segments"]):
            secs = seg["seconds"] if seg["seconds"] is not None else scene.seconds - used
            used += secs
            out.append((f"segments[{i}]", seg["clip"], secs))
        return out

    def walk(value: Any, path: str) -> None:
        if isinstance(value, dict):
            if "video" in value and "rate" in value:
                out.append((path, value, scene.seconds))
                return
            for k, v in value.items():
                if k in ("step_at", "legend"):  # times, not played clips
                    continue
                walk(v, f"{path}.{k}" if path else k)
        elif isinstance(value, list):
            for i, v in enumerate(value):
                walk(v, f"{path}[{i}]")

    walk(raw, "")
    return out


def overrun(demo: Demo, folder: Path, cut: Cut, *, tolerance_s: float = 0.0) -> Result:
    """Every played clip stays inside its beat: ``offset + seconds × rate ≤ beat end``.

    A clip ref without a beat is checked against its media file's length;
    when that file can't be read the check is ``unknown``.
    """
    marks = load_marks(folder, demo, cut.lang)
    files = clip_files(demo, cut.lang)
    fails, unknown = [], []
    for scene in demo.scenes:
        for where, ref, secs in _played(scene):
            spec = demo.clips[ref["video"]]
            used = ref["offset"] + secs * ref["rate"]
            if ref.get("beat"):
                start, end = marks[spec.marks or ref["video"]][ref["beat"]]
                limit, what = end - start, f"beat {ref['beat']!r}"
            else:
                path = media_root(folder, demo) / files[ref["video"]]
                try:
                    limit = probe(path)["duration"]
                except Exception:  # noqa: BLE001 — any probe failure is "can't tell", reported as unknown
                    unknown.append(f"{scene.id}.{where}: can't read {files[ref['video']]}")
                    continue
                limit, what = limit, f"video {ref['video']!r}"
                if limit <= 0:
                    unknown.append(f"{scene.id}.{where}: {files[ref['video']]} has no length")
                    continue
            if used > limit + tolerance_s:
                fails.append(f"{scene.id}.{where}: {ref['offset']:g} s + {secs:g} s × {ref['rate']:g} = {used:.1f} s "
                             f"> {what} {limit:.1f} s (overruns {used - limit:.1f} s)")
    name = f"overrun ({cut.id})"
    if fails:
        return Result(name, "fail", "; ".join(fails))
    if unknown:
        return Result(name, "unknown", "; ".join(unknown))
    return Result(name, "pass")


# ----------------------------------------------------------------------------- timer / state timeline


def _hex(c: str) -> np.ndarray:
    c = c.lstrip("#")
    return np.array([int(c[i:i + 2], 16) for i in (0, 2, 4)], dtype=np.float64)


def classify_frame(pixels: np.ndarray, palette: list[np.ndarray], *, max_dist: float = 90.0) -> Optional[int]:
    """The palette index most pixels of a frame are nearest to, ignoring the pale background; None if unsure."""
    px = pixels.reshape(-1, 3).astype(np.float64)
    pale = (px.min(axis=1) > 200) & (px.max(axis=1) - px.min(axis=1) < 30)
    px = px[~pale]
    if len(px) < 20:
        return None
    dists = np.stack([np.linalg.norm(px - p, axis=1) for p in palette], axis=1)
    nearest = dists.argmin(axis=1)
    close = dists.min(axis=1) < max_dist
    if close.sum() < 20:
        return None
    counts = np.bincount(nearest[close], minlength=len(palette))
    return int(counts.argmax())


def state_timeline(clip: Path, *, box: dict, palette: list[str], start: float, end: float,
                   step_s: float = 0.5) -> list[list[float]]:
    """``[[seconds after start, state], …]`` — one entry per change of the dominant palette colour in ``box``."""
    w, h = 64, max(1, round(64 * box["h"] / box["w"]))
    raw = ffmpeg_pipe(["-ss", f"{start:.3f}", "-t", f"{end - start:.3f}", "-i", str(clip),
                       "-vf", f"fps={1 / step_s:g},crop={int(box['w'])}:{int(box['h'])}:{int(box['x'])}:{int(box['y'])},"
                              f"scale={w}:{h}",
                       "-f", "rawvideo", "-pix_fmt", "rgb24", "-"])
    frame = w * h * 3
    pal = [_hex(c) for c in palette]
    out: list[list[float]] = []
    last: Optional[int] = None
    for i in range(len(raw) // frame):
        px = np.frombuffer(raw[i * frame:(i + 1) * frame], dtype=np.uint8).reshape(h, w, 3)
        state = classify_frame(px, pal)
        if state is not None and state != last:
            out.append([round(i * step_s, 2), state])
            last = state
    return out


# ----------------------------------------------------------------------------- privacy


def fold(text: str) -> str:
    """Lowercase, accents stripped: 'Andrés' and 'andres' compare equal."""
    return "".join(c for c in unicodedata.normalize("NFKD", text.lower()) if not unicodedata.combining(c))


def tokens(text: str) -> set[str]:
    return {t for t in re.findall(r"\w+", fold(text)) if len(t) > 3}


def read_roster(path: Path) -> list[str]:
    """Names from a .txt/.csv (first column, one per line) or an .xlsx with a "name" column."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook  # noqa: PLC0415

        ws = load_workbook(path, read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(c or "").strip().lower() for c in rows[0]]
        col = header.index("name") if "name" in header else 0
        return [str(r[col]).strip() for r in rows[1:] if r and r[col]]
    with path.open(encoding="utf-8-sig", newline="") as f:
        return [row[0].strip() for row in csv.reader(f) if row and row[0].strip() and not row[0].startswith("#")]


def _strings(value: Any) -> Iterable[str]:
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for v in value.values():
            yield from _strings(v)
    elif isinstance(value, list):
        for v in value:
            yield from _strings(v)


def privacy(demo: Demo, folder: Path, *, roster: Optional[list[str]] = None) -> Result:
    """No real name — full name or any name token over 3 letters — anywhere a cut can show text;
    no blocklisted term in the copy of a public cut's language or in the scenes."""
    cfg = demo.privacy
    if roster is None and cfg.roster:
        path = folder / cfg.roster
        if not path.is_file():
            return Result("privacy", "unknown", f"roster file not found: {cfg.roster}")
        roster = read_roster(path)
    all_text = list(_strings(demo.copy_)) + list(_strings([s.model_dump(mode="json") for s in demo.scenes]))
    all_text += list(_strings(demo.product.model_dump(mode="json")))
    for extra in cfg.scan:
        p = folder / extra
        if not p.is_file():
            return Result("privacy", "unknown", f"scan file not found: {extra}")
        all_text.append(p.read_text(encoding="utf-8", errors="replace"))
    shown = set().union(*(tokens(t) for t in all_text)) if all_text else set()
    folded = "\n".join(fold(t) for t in all_text)
    hits = []
    for name in roster or []:
        if fold(name) in folded:
            hits.append(f"real name {name!r} appears")
            continue
        shared = sorted((tokens(name) & shown) - {fold(a) for a in cfg.allow})
        if shared:
            hits.append(f"{name!r} shares {', '.join(repr(t) for t in shared)}")
    public_langs = {c.lang for c in demo.cuts if c.public}
    public_text = [t for lang in public_langs for t in _strings(demo.copy_.get(lang, {}))]
    public_text += list(_strings([s.model_dump(mode="json") for s in demo.scenes]))
    public_folded = "\n".join(fold(t) for t in public_text)
    for term in cfg.blocklist:
        if public_langs and re.search(rf"(?<!\w){re.escape(fold(term))}(?!\w)", public_folded):
            hits.append(f"blocklisted term {term!r} in a public cut")
    if hits:
        return Result("privacy", "fail", "; ".join(hits))
    if not roster and not cfg.blocklist:
        return Result("privacy", "pass", "no roster or blocklist declared — nothing to compare against")
    return Result("privacy", "pass", f"{len(roster or [])} roster names, {len(cfg.blocklist)} blocklist terms")


# ----------------------------------------------------------------------------- licence


def licence(demo: Demo) -> Result:
    """A public cut never carries a private-only track."""
    bad = [f"{c.id} uses {t.file}" for c in demo.cuts if c.public for t in c.soundtrack if t.licence == "private-only"]
    if bad:
        return Result("licence", "fail", "private-only music in a public cut: " + "; ".join(bad))
    return Result("licence", "pass")


# ----------------------------------------------------------------------------- all


def run_checks(demo: Demo, folder: Path, cuts: list[Cut]) -> list[Result]:
    results = [overrun(demo, folder, cut) for cut in cuts]
    results += [privacy(demo, folder), licence(demo)]
    return results


def write_report(folder: Path, results: list[Result]) -> Path:
    path = folder / "out" / "prep" / "checks.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps([r.__dict__ for r in results], ensure_ascii=False, indent=1), encoding="utf-8")
    return path
